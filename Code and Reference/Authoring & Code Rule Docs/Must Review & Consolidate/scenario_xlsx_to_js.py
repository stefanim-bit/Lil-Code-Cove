#!/usr/bin/env python3
"""
Converts a filled-in Practice_Session_Scenario_Template.xlsx into a
`window.SCENARIO = { ... }` JS block ready to paste into
practice_scenario_template_v2.html's SCENARIO <script> tag.

Usage:
    python3 scenario_xlsx_to_js.py <input.xlsx> [output.js]

If output.js is omitted, prints the JS block to stdout.
"""
import sys
import json
import openpyxl


def rows(ws):
    """Yield each data row (after the header) as a list of cell values, skipping blank rows."""
    it = ws.iter_rows(min_row=2, values_only=True)
    for r in it:
        if any(v not in (None, "") for v in r):
            yield r


def s(v):
    return "" if v is None else str(v).strip()


def js_str(v):
    return json.dumps(s(v), ensure_ascii=False)


def load(path):
    return openpyxl.load_workbook(path, data_only=True)


def build_meta(wb):
    meta = {}
    for field, value, _notes in rows(wb["Meta"]):
        meta[s(field)] = s(value)
    return meta


def build_stages(wb):
    out = []
    total = 0
    for id_, name, max_, strength, improvement in rows(wb["Stages"]):
        if id_ is None or max_ is None:
            continue
        max_i = int(max_)
        total += max_i
        out.append(
            "    { id:%s, name:%s, max:%d,\n"
            "      strength:%s,\n"
            "      improvement:%s }"
            % (js_str(id_), js_str(name), max_i, js_str(strength), js_str(improvement))
        )
    if total != 100:
        print(f"WARNING: Stages max values sum to {total}, not 100.", file=sys.stderr)
    return ",\n".join(out)


def build_screens_and_choices(wb):
    screens = {}
    for screen_id, stage, gate, learner, speaker, prompt in rows(wb["Screens"]):
        if screen_id is None:
            continue
        screens[s(screen_id)] = {
            "stage": s(stage),
            "gate": s(gate).upper().startswith("Y"),
            "learnerLine": s(learner),
            "speaker": s(speaker),
            "prompt": s(prompt),
            "choices": [],
        }

    for row in rows(wb["Choices"]):
        (screen_id, choice_id, text, verdict, points, goTo, retry, capIfRetried,
         r1_speaker, r1_text, r2_speaker, r2_text, r3_speaker, r3_text,
         stageOverride, credit_stage, credit_points, credit_verdict) = (
            list(row) + [None] * (18 - len(row))
        )[:18]
        if screen_id is None:
            continue
        sid = s(screen_id)
        if sid not in screens:
            raise ValueError(f"Choices row references unknown screen_id '{sid}' — add it to the Screens tab.")
        # choice_id (column B) is only for readability inside the sheet — the
        # SCENARIO schema has no `id` field on choices, so it is not emitted below.
        reactions = []
        for speaker, text_ in ((r1_speaker, r1_text), (r2_speaker, r2_text), (r3_speaker, r3_text)):
            if s(text_):
                reactions.append({"speaker": s(speaker) or "caller", "text": s(text_)})
        choice = {
            "text": s(text),
            "verdict": s(verdict),
            "points": int(points) if points not in (None, "") else 0,
            "goTo": s(goTo),
            "retry": s(retry).upper().startswith("Y"),
            "capIfRetried": capIfRetried,
            "reactions": reactions,
            "stageOverride": s(stageOverride),
            "credit_stage": s(credit_stage),
            "credit_points": credit_points,
            "credit_verdict": s(credit_verdict),
        }
        screens[sid]["choices"].append(choice)

    return screens


def render_choice(ch, indent="        "):
    parts = [f'{{ text:{js_str(ch["text"])},']
    parts.append(f'  verdict:{js_str(ch["verdict"])}, points:{ch["points"]}, goTo:{js_str(ch["goTo"])}')
    if ch["retry"]:
        cap = ch["capIfRetried"] if ch["capIfRetried"] not in (None, "") else 0
        parts.append(f', retry:true, capIfRetried:{int(cap)}')
    if ch["stageOverride"]:
        parts.append(f', stageOverride:{js_str(ch["stageOverride"])}')
    if ch["reactions"]:
        items = ", ".join(
            f'{{ speaker:{js_str(r["speaker"])}, text:{js_str(r["text"])} }}' for r in ch["reactions"]
        )
        parts.append(f',\n{indent}  reactions:[{items}]')
    if ch["credit_stage"]:
        cp = int(ch["credit_points"]) if ch["credit_points"] not in (None, "") else 0
        cv = ch["credit_verdict"] or "correct"
        parts.append(
            f',\n{indent}  credits:[{{ stage:{js_str(ch["credit_stage"])}, points:{cp}, verdict:{js_str(cv)} }}]'
        )
    parts.append(" }")
    return "".join(parts)


def render_screens(screens, order):
    blocks = []
    for sid in order:
        scr = screens[sid]
        head = f'    {sid}: {{\n      stage:{js_str(scr["stage"])}, gate:{"true" if scr["gate"] else "false"},'
        lines = [head]
        if scr.get("learnerLine"):
            lines.append(f'      learnerLine:{js_str(scr["learnerLine"])},')
        if scr["speaker"]:
            lines.append(f'      speaker:{js_str(scr["speaker"])},')
        lines.append(f'      prompt:{js_str(scr["prompt"])},')
        lines.append("      choices:[")
        choice_lines = [f"        {render_choice(ch)}" for ch in scr["choices"]]
        lines.append(",\n".join(choice_lines))
        lines.append("      ]")
        lines.append("    }")
        blocks.append("\n".join(lines))
    return ",\n\n".join(blocks)


def build_resources(wb):
    """Reads the single 'Resources' tab (title/url/points/note) into the
    `resources: { items: [...] }` shape used by the merged Resources panel
    (see SCENARIO_AUTHORING_HANDOFF.md §7b). This replaced the old separate
    Open CSI / Ask a Colleague / Check the Articles helpers."""
    if "Resources" not in wb.sheetnames:
        return None

    item_lines = []
    for title, url, points, note in rows(wb["Resources"]):
        if not s(title):
            continue
        pts = int(points) if points not in (None, "") else 0
        line = f'    {{ title:{js_str(title)}, url:{js_str(url)}, points:{pts}'
        if s(note):
            line += f', note:{js_str(note)}'
        line += " }"
        item_lines.append(line)

    if not item_lines:
        return None

    items_joined = ",\n".join(item_lines)
    return f'''  resources: {{
    items: [
{items_joined}
    ]
  }},
'''


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    in_path = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) > 2 else None

    wb = load(in_path)
    meta = build_meta(wb)
    stages_js = build_stages(wb)
    screens = build_screens_and_choices(wb)
    screen_order = list(screens.keys())
    screens_js = render_screens(screens, screen_order)
    resources_block = build_resources(wb) or ""

    out = f"""window.SCENARIO = {{
  title: {js_str(meta.get('title'))},
  userBio: {js_str(meta.get('userBio'))},
{resources_block}
  stages: [
{stages_js}
  ],

  start: {js_str(meta.get('start'))},
  screens: {{

{screens_js}

  }}
}};
"""

    if out_path:
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(out)
        print(f"Wrote {out_path}")
    else:
        print(out)


if __name__ == "__main__":
    main()
